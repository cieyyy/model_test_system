import asyncio
import importlib.util
import json
import math
from pathlib import Path
import tempfile
import time
import unittest

spec = importlib.util.spec_from_file_location("acceptance", Path(__file__).resolve().parents[1] / "run_acceptance.py")
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)


class MetricsTests(unittest.TestCase):
    def test_existing_command_style_aliases(self):
        args, config = m.settings(["--base-url", "https://example.com/v1", "--key-env", "san_API_KEY",
                                   "--model", "kimi-k3", "--run", "--confirm-paid"])
        self.assertTrue(args.execute)
        self.assertTrue(args.confirm_paid)
        self.assertEqual(config["api_key_env"], "san_API_KEY")
        self.assertEqual(config["model"], "kimi-k3")

    def test_decimal_timing_options(self):
        _, c = m.settings(["--timeout-seconds", "90", "--duration-seconds", "60.5", "--p50-ms", "2500.5"])
        self.assertEqual(c["timeout_seconds"], 90)
        self.assertEqual(c["duration_seconds"], 60.5)
        self.assertEqual(c["p50_ms"], 2500.5)

    def records(self):
        return [dict(case_id="real_traffic", result="pass", usage_valid=True, cache_valid=True,
                     input_tokens=1_000_000, output_tokens=10, total_tokens=1_000_010,
                     cached_tokens=850_000, ttft_ms=2500, elapsed_ms=3000,
                     modalities=["image" if i % 2 else "video"])
                for i in range(100)]

    def summary(self, rows, elapsed=180):
        c = dict(m.DEFAULT, cases=["real_traffic"])
        return m.summarize(rows, c, {}, 0, elapsed)["real_traffic"]

    def test_simultaneous_thresholds(self):
        r = self.summary(self.records())
        self.assertEqual(r["result"], "pass")
        self.assertEqual(r["detail"]["cache_ratio"], .85)

    def test_drain_time_prevents_false_tpm_pass(self):
        r = self.summary(self.records(), 1000)
        self.assertEqual(r["result"], "fail")
        self.assertFalse(r["detail"]["checks"]["achieved_tpm"])

    def test_missing_cache_not_excluded(self):
        rows = self.records()
        rows[0].update(cache_valid=False, cached_tokens=None)
        self.assertFalse(self.summary(rows)["detail"]["checks"]["cache_85_percent"])

    def test_failures_in_percentile(self):
        rows = self.records()
        for r in rows[:11]:
            r["result"] = "error"
        report = self.summary(rows)
        self.assertIsNone(report["detail"]["p90_ms"])
        self.assertFalse(report["detail"]["checks"]["p90"])
        self.assertEqual(m.percentile([1, math.inf], .5), None)

    def test_real_video_required(self):
        rows = self.records()
        for r in rows:
            r["modalities"] = ["image"]
        self.assertFalse(self.summary(rows)["detail"]["checks"]["mixed_image_video"])

    def test_invalid_usage(self):
        self.assertFalse(m.usage_fields(["malformed"])["usage_valid"])
        self.assertFalse(m.usage_fields(dict(prompt_tokens=10, completion_tokens=2, total_tokens=15))["usage_valid"])
        self.assertFalse(m.usage_fields(dict(prompt_tokens=10, completion_tokens=2, total_tokens=12,
                                           prompt_tokens_details=dict(cached_tokens=11)))["cache_valid"])

    def test_context_minimum_uses_provider_tokens_without_artificial_ceiling(self):
        self.assertTrue(m.context_length_ok(1_019_621, m.DEFAULT))
        self.assertTrue(m.context_length_ok(1_000_000, m.DEFAULT))
        for count in (999_999, None, "1000000", True):
            self.assertFalse(m.context_length_ok(count, m.DEFAULT))

    def test_quota_error_is_distinct_from_invalid_credentials(self):
        parsed = m.provider_error('{"error":{"message":"Not enough quota to pre-consume for this request"}}', 403)
        self.assertEqual(parsed["category"], "insufficient_quota")

    def test_provider_cache_field_fallbacks(self):
        base = dict(prompt_tokens=100, completion_tokens=20, total_tokens=120)
        for fields, expected in (
            (dict(cached_tokens=85), 85),
            (dict(prompt_cache_hit_tokens=86), 86),
            (dict(prompt_tokens_details=dict(cached_tokens=0), cached_tokens=85), 0),
        ):
            parsed = m.usage_fields(dict(base, **fields))
            self.assertTrue(parsed["cache_valid"])
            self.assertEqual(parsed["cached_tokens"], expected)
        self.assertFalse(m.usage_fields(base)["cache_valid"])

    def test_100mb_actual_bytes(self):
        row = dict(body=dict(messages=[dict(role="user", content="只回复 BODY-OK")]))
        raw = m.payload("body_100mb", row, m.DEFAULT)
        self.assertEqual(len(raw), 100_000_000)
        self.assertEqual(json.loads(raw)["model"], m.DEFAULT["model"])

    def test_no_budget_truncation(self):
        with self.assertRaises(ValueError):
            m.plan(dict(m.DEFAULT, cases=["body_100mb"], max_requests=2))

    def test_deployment_not_blackbox_pass(self):
        report = m.summarize([], dict(m.DEFAULT, cases=["deployment"]), {})
        self.assertEqual(report["deployment"]["result"], "manual_review")

    def test_synthetic_scope_is_explicit(self):
        rows = [dict(r, case_id="synthetic_traffic") for r in self.records()]
        report = m.summarize(rows, dict(m.DEFAULT, cases=["synthetic_traffic"]), {}, 0, 180)
        self.assertNotIn("real_traffic", report)
        self.assertEqual(report["synthetic_traffic"]["detail"]["business_cache_acceptance"], "not_evaluated")

    def test_report_actual_values_below_threshold_and_partial_cache(self):
        c = dict(m.DEFAULT, cases=["real_traffic"])
        rows = self.records()
        for r in rows:
            r.update(ttft_ms=7000, cached_tokens=400_000)
        rows[0].update(cache_valid=False, cached_tokens=None)
        rows[1].update(result="error", ttft_ms=None)
        summary = m.summarize(rows, c, {}, 0, 1000)
        metrics = {row["metric"]: row for row in m.measurements(rows, summary, c)}
        self.assertEqual(metrics["观测 TTFT P90"]["value"], 7000)
        self.assertEqual(metrics["观测 TTFT P90"]["samples"], 99)
        self.assertEqual(metrics["缓存命中率"]["value"], 40)
        self.assertEqual(metrics["缓存命中率"]["status"], "部分样本")
        self.assertEqual(metrics["缓存命中率"]["samples"], 99)
        self.assertEqual(metrics["Input TPM"]["value"], 5_940_000)
        self.assertEqual(metrics["成功率"]["value"], 99)

    def test_missing_values_not_zero_or_false(self):
        c = dict(m.DEFAULT, cases=["deployment"])
        summary = m.summarize([], c, {})
        self.assertIsNone(summary["deployment"]["detail"]["fp8_declared"])
        metrics = m.measurements([], summary, c)
        self.assertTrue(all(r["value"] is None and r["status"] == "未提供" for r in metrics))

    def test_fixture_expansion_and_path_safety(self):
        import base64
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root / "image.jpg").write_bytes(b"test JPEG bytes")
            (root / "prefix.txt").write_text("检索资料", encoding="utf-8")
            c = dict(m.DEFAULT, asset_dir=td)
            image = m.resolve_fixtures("fixture://image.jpg", c)
            self.assertEqual(base64.b64decode(image.split(",", 1)[1]), b"test JPEG bytes")
            self.assertEqual(m.resolve_fixtures("fixture-text://prefix.txt", c), "检索资料")
            with self.assertRaises(ValueError):
                m.fixture_path("../escape.txt", c)

    def test_public_materials_are_runnable(self):
        # 随交付附带的真实下载素材：不访问外网，不调用模型。
        root = Path(__file__).resolve().parents[1]
        preset = root / "config.public.json"
        if not preset.is_file():
            self.skipTest("尚未准备公开素材")
        _, c = m.settings(["--config", str(preset)])
        tasks, missing, _ = m.plan(c)
        self.assertFalse(missing)
        self.assertEqual(len(tasks), 433)
        raw = m.payload(tasks[0][0], tasks[0][1], c)
        self.assertEqual(len(raw), 100_000_000)
        self.assertIsInstance(json.loads(raw)["messages"][0]["content"], str)
        media_row = m.read_json(c["body_file"])
        raw = m.payload("body_100mb", media_row, dict(c, body_mode="media"))
        self.assertEqual(len(raw), 100_000_000)
        body = json.loads(raw)
        content = body["messages"][-1]["content"]
        self.assertGreater(sum(p.get("type") == "image_url" for p in content), 100)
        self.assertTrue(any(p.get("type") == "video_url" for p in content))
        raw_json_bytes = len(m.dumps(body).encode("utf-8"))
        self.assertGreater(raw_json_bytes, 99_800_000)
        context = next(row for case, row, _ in tasks if case == "context_1m")
        ctx = json.loads(m.payload("context_1m", context, c))["messages"][0]["content"]
        self.assertTrue(all(answer in ctx for answer in context["expected_contains"]))

    def test_public_samples_cannot_masquerade_as_real(self):
        rows = [dict(offset_seconds=i, fixture_type="public_synthetic",
                     body={"messages": [{"role": "user", "content": "test"}]}) for i in range(100)]
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "traffic.jsonl"
            path.write_text("\n".join(json.dumps(r) for r in rows), encoding="utf-8")
            with self.assertRaises(ValueError):
                m.plan(dict(m.DEFAULT, cases=["real_traffic"], traffic_file=str(path)))

    def test_report_export(self):
        from openpyxl import load_workbook
        c = dict(m.DEFAULT, cases=["real_traffic", "deployment"])
        rows = self.records()
        report = m.summarize(rows, c, {}, 0, 180)
        with tempfile.TemporaryDirectory() as td:
            m.export(Path(td), c, rows, report)
            wb = load_workbook(Path(td) / "results.xlsx")
            self.assertEqual(wb.active.title, "实测指标")
            ws = wb.worksheets[1]
            self.assertEqual(ws["H5"].value, "通过")
            self.assertEqual(ws["H6"].value, "待人工核验")
            self.assertEqual(ws.auto_filter.ref, "A4:J6")
            self.assertEqual(ws.freeze_panes, "A5")
            self.assertTrue((Path(td) / "requests.csv").is_file())
            self.assertTrue((Path(td) / "metrics.csv").is_file())
            cells = list(wb["实测指标"].values)
            cache = next(row for row in cells if row[1] == "缓存命中率")
            self.assertEqual(cache[2], 85)


class StreamingTests(unittest.IsolatedAsyncioTestCase):
    async def test_utf8_fragmentation_real_http(self):
        # 本地回环服务验证真实逐字节网络分片，绝不访问付费端点。
        import aiohttp
        from aiohttp import web
        async def handle(req):
            await req.read()
            response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
            await response.prepare(req)
            objects = [dict(model="test", choices=[dict(index=0, delta=dict(role="assistant"))]),
                       dict(model="test", choices=[dict(index=0, delta=dict(content="中文答案"))]),
                       dict(model="test", choices=[dict(index=0, delta={}, finish_reason="stop")],
                            usage=dict(prompt_tokens=100, completion_tokens=4, total_tokens=104,
                                       prompt_tokens_details=dict(cached_tokens=85)))]
            raw = "".join("data: " + json.dumps(o, ensure_ascii=False) + "\n\n" for o in objects) + "data: [DONE]\n\n"
            for b in raw.encode():
                await response.write(bytes([b]))
            await response.write_eof()
            return response
        app = web.Application()
        app.router.add_post("/v1/chat/completions", handle)
        runner = web.AppRunner(app)
        await runner.setup()
        site = web.TCPSite(runner, "127.0.0.1", 0)
        await site.start()
        port = site._server.sockets[0].getsockname()[1]
        try:
            with tempfile.TemporaryDirectory() as td:
                async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=5)) as session:
                    row = dict(body=dict(messages=[dict(role="user", content="问题")]), expected_contains=["中文答案"])
                    r = await m.request(session, asyncio.Semaphore(1), asyncio.Event(), "test-only-key",
                                        dict(m.DEFAULT, base_url=f"http://127.0.0.1:{port}/v1"),
                                        "real_traffic", row, 1, time.perf_counter(), Path(td))
                self.assertEqual(r["result"], "pass")
                self.assertEqual(r["content"], "中文答案")
                self.assertIsNotNone(r["ttft_ms"])
                self.assertEqual(r["cached_tokens"], 85)
                evidence = json.loads((Path(td) / "000001.json").read_text())
                self.assertEqual(len(evidence["events"]), 4)
        finally:
            await runner.cleanup()


class FailureIsolationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        import aiohttp
        from aiohttp import web
        self.temp = tempfile.TemporaryDirectory()
        self.stop = asyncio.Event()
        self.seen = []
        async def handle(req):
            data = await req.json()
            text = data["messages"][0]["content"]
            self.seen.append(text)
            if text == "bad400":
                return web.Response(status=400, content_type="text/event-stream",
                    text='data: {"error":{"message":"Unsupported video_url. key=test-only-key"}}\n\n')
            if text == "bad401":
                return web.json_response({"error": {"message": "invalid key test-only-key"}}, status=401)
            if text == "sse_error":
                return web.Response(content_type="text/event-stream",
                                    text='data: {"error":{"message":"context length exceeded"}}\n\n')
            if text == "huge_error":
                return web.Response(status=413, text="payload too large " + "x" * 100000)
            return web.Response(content_type="text/event-stream", text='data: ' + json.dumps(dict(
                model="test", choices=[dict(index=0, delta=dict(content="OK"), finish_reason="stop")],
                usage=dict(prompt_tokens=100, completion_tokens=1, total_tokens=101,
                           prompt_tokens_details=dict(cached_tokens=85)))) + '\n\ndata: [DONE]\n\n')
        app = web.Application()
        app.router.add_post("/v1/chat/completions", handle)
        self.runner = web.AppRunner(app)
        await self.runner.setup()
        site = web.TCPSite(self.runner, "127.0.0.1", 0)
        await site.start()
        port = site._server.sockets[0].getsockname()[1]
        self.c = dict(m.DEFAULT, base_url=f"http://127.0.0.1:{port}/v1")
        self.session = aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=5))

    async def asyncTearDown(self):
        await self.session.close()
        await self.runner.cleanup()
        self.temp.cleanup()

    async def send(self, text, index, case_stop=None, due=None):
        return await m.request(self.session, asyncio.Semaphore(1), self.stop, "test-only-key", self.c,
            "diagnostic", dict(body={"messages": [{"role": "user", "content": text}]}, expected_contains=["OK"]),
            index, due or time.perf_counter(), Path(self.temp.name), case_stop)

    async def test_400_body_saved_and_next_case_continues(self):
        case_stop = asyncio.Event()
        first = await self.send("bad400", 1, case_stop)
        self.assertFalse(self.stop.is_set())
        self.assertTrue(case_stop.is_set())
        self.assertEqual(first["provider_error"]["category"], "video_format_or_support")
        self.assertIn("Unsupported video_url", first["error_body"])
        self.assertNotIn("test-only-key", (Path(self.temp.name) / "000001.json").read_text())
        skipped = await self.send("OK", 2, case_stop)
        self.assertFalse(skipped["attempted"])
        self.assertEqual(skipped["result"], "skipped")
        next_case = await self.send("OK", 3, asyncio.Event())
        self.assertEqual(next_case["result"], "pass")
        self.assertEqual(self.seen, ["bad400", "OK"])

    async def test_auth_failure_globally_stops_without_waiting_for_schedule(self):
        first = await self.send("bad401", 1)
        self.assertEqual(first["provider_error"]["category"], "authentication_or_permission")
        self.assertTrue(self.stop.is_set())
        skipped = await asyncio.wait_for(self.send("OK", 2, due=time.perf_counter() + 60), timeout=1)
        self.assertEqual(skipped["error_type"], "auth_stopped")
        self.assertEqual(self.seen, ["bad401"])

    async def test_sse_200_error_preserves_message(self):
        record = await self.send("sse_error", 1)
        self.assertEqual(record["result"], "error")
        self.assertEqual(record["provider_error"]["category"], "context_limit")
        self.assertIn("context length exceeded", record["error_body"])

    async def test_error_body_bounded(self):
        record = await self.send("huge_error", 1)
        self.assertEqual(record["provider_error"]["category"], "body_size_limit")
        self.assertTrue(record["error_body_truncated"])
        self.assertEqual(len(record["error_body"]), 65536)

    async def test_skipped_is_not_model_latency(self):
        first = await self.send("OK", 1)
        stop = asyncio.Event()
        stop.set()
        skipped = await self.send("OK", 2, stop)
        summary = m.summarize([first, skipped], dict(self.c, cases=["diagnostic"]), {})
        metrics = {r["metric"]: r for r in m.measurements([first, skipped], summary, self.c)}
        self.assertEqual(metrics["实际发送请求数"]["value"], 1)
        self.assertEqual(metrics["跳过请求数"]["value"], 1)
        self.assertEqual(metrics["成功率"]["value"], 100)
        self.assertEqual(metrics["单请求总耗时 P50"]["samples"], 1)


if __name__ == "__main__":
    unittest.main()
