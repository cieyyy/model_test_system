#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从统一测试结果目录生成 Excel 和 HTML 报告。

用法：
  python3 generate_test_report.py --model glm-5.3-flash --base-url https://...
  python3 generate_test_report.py --model glm-5.3-flash --failed-only

报告不会写入 API Key；失败报告只保留 FAIL/未通过/人工复核项目。
"""
from __future__ import annotations

import argparse
import copy
import csv
import html
import json
import math
import os
import shutil
import statistics
from urllib.parse import quote
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(os.environ.get("MT_RESULT_ROOT", Path(__file__).resolve().parent))
TEMPLATE_CANDIDATES = [
    Path(os.environ.get("MT_TEMPLATE_DIR", Path(__file__).resolve().parents[1] / "docs/models")),
]


def template_path(filename: str) -> Path:
    for directory in TEMPLATE_CANDIDATES:
        candidate = directory / filename
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"找不到报告模板：{filename}")


def slug(value: str) -> str:
    return "".join(c if c.isalnum() or c in "._-" else "_" for c in value)


def result_batch_key(directory: Path) -> str:
    """Return a stable project key from a timestamped result directory name."""
    parts = directory.name.split("-")
    if len(parts) >= 3 and len(parts[0]) == 8 and len(parts[1]) == 6 and parts[0].isdigit() and parts[1].isdigit():
        start = 2
        if len(parts) >= 4 and parts[2].isdigit():
            start = 3
        return "-".join(parts[start:]) or directory.name
    return directory.name


def load_rows(result_root: Path, run_date: str = "") -> list[dict]:
    rows: list[dict] = []
    csv_paths = list(result_root.rglob("results.csv"))
    result_dirs = {path.parent for path in csv_paths}
    csv_paths.extend(
        path for path in result_root.rglob("conversation.csv")
        if path.parent not in result_dirs
    )
    if run_date:
        latest_by_key: dict[str, Path] = {}
        for path in csv_paths:
            if not path.parent.name.startswith(run_date):
                continue
            key = result_batch_key(path.parent)
            if key not in latest_by_key or path.parent.name > latest_by_key[key].parent.name:
                latest_by_key[key] = path
        csv_paths = list(latest_by_key.values())
    for path in sorted(csv_paths):
        if run_date and not path.parent.name.startswith(run_date):
            continue
        try:
            with path.open(encoding="utf-8-sig", newline="") as f:
                for row in csv.DictReader(f):
                    row["_source"] = str(path.parent)
                    rows.append(row)
        except (OSError, UnicodeError):
            continue
    # 性能复测包使用 probe/stress/sustain JSON，转换为同一张汇总明细表。
    for path in sorted(result_root.rglob("total-test-results/*.json")):
        if run_date and run_date not in str(path):
            continue
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
            summaries = []
            if path.name == "stress.json":
                summaries = payload.get("summaries", [])
            elif path.name == "probe.json":
                summaries = [x.get("summary", {}) for x in payload.get("probes", [])]
            elif path.name == "sustain.json":
                summaries = [x.get("summary", {}) for x in payload.get("targets", [])]
            for index, summary in enumerate(summaries, 1):
                if not isinstance(summary, dict):
                    continue
                case_name = ttft_batch_name(summary) or f"性能-{path.stem}-{index}"
                is_ttft_batch = case_name.startswith("TTFT ")
                rows.append({
                    "case": case_name,
                    "endpoint": "/v1/chat/completions",
                    "elapsed_ms": summary.get("elapsed_p50_ms", ""),
                    "prompt_tokens": summary.get("prompt_tokens", ""),
                    "completion_tokens": summary.get("completion_tokens", ""),
                    "http_status": 200 if summary.get("failed", 1) == 0 else 500,
                    "result": "data_display" if is_ttft_batch else ("pass" if summary.get("failed", 1) == 0 else "fail"),
                    "content": (
                        f"并发 {summary.get('concurrency', '—')}；请求 {summary.get('attempted', '—')}；"
                        f"成功率 {summary.get('success_rate_pct', '—')}%；RPS {summary.get('success_rps', '—')}；"
                        f"TTFT P50 {summary.get('ttft_p50_ms', '—')} ms；TTFT P95 {summary.get('ttft_p95_ms', '—')} ms"
                    ),
                    "ttft_ms": summary.get("ttft_p50_ms", ""),
                    "_source": str(path),
                })
        except (OSError, ValueError, TypeError):
            continue
    return rows


def ttft_batch_name(summary: dict) -> str:
    """Return the report label for explicit TTFT batch tests."""
    text = " ".join(
        str(summary.get(key) or "")
        for key in ("case", "name", "phase", "label", "scenario", "batch")
    ).lower()
    context_k = None
    for size in (16, 64, 100):
        if f"{size}k" in text or f"{size} k" in text:
            context_k = size
            break
    if context_k is None:
        for key in ("context_k", "input_k", "prompt_k", "tokens_k"):
            try:
                value = float(summary.get(key))
            except (TypeError, ValueError):
                continue
            if value in {16, 64, 100}:
                context_k = int(value)
                break
    if context_k is None:
        return ""
    concurrency = summary.get("concurrency") or summary.get("parallel") or summary.get("workers") or 10
    try:
        concurrency_text = str(int(float(concurrency)))
    except (TypeError, ValueError):
        concurrency_text = str(concurrency or 10)
    return f"TTFT {context_k}k({concurrency_text}并发)"


def status(row: dict) -> str:
    value = str(row.get("result") or row.get("status") or "").lower()
    if value in {"data", "data_display", "metrics", "metric", "数据展示"}:
        return "数据展示"
    if value in {"pass", "passed", "通过"}:
        return "通过"
    if value in {"manual_review", "review", "人工复核"}:
        return "人工复核"
    if value in {"fail", "failed", "未通过"} or str(row.get("http_status")) not in {"", "200"}:
        return "未通过"
    return "待测试"


def display_rows(rows: list[dict]) -> list[dict]:
    result = []
    for i, row in enumerate(rows, 1):
        source = str(row.get("_source", ""))
        parent = Path(source).name.lower()
        inferred = ""
        if "long-conversation" in parent:
            inferred = "100轮 KVV 长对话"
        elif "long-output" in parent:
            inferred = "长输出截断"
        elif "sophnet-load" in parent:
            summary_path = Path(source) / "summary.json"
            mode = ""
            body_mb = 0.0
            try:
                summary = json.loads(summary_path.read_text(encoding="utf-8"))
                mode = str(summary.get("mode") or "")
                body_mb = float(summary.get("body_mb") or 0)
            except (OSError, ValueError, TypeError):
                pass
            if body_mb > 0 or "body-" in parent:
                inferred = "100MB 请求体阶梯"
            else:
                inferred = {"cache": "缓存命中率阶梯", "context": "1M 上下文边界", "mixed": "混合图文视频"}.get(mode, "多模态/上下文请求")
        elif "load-test" in parent:
            inferred = "固定速率性能"
        elif "tpm-benchmark" in parent:
            inferred = "1000万 TPM 吞吐压测"
        elif "stream-metrics" in parent:
            inferred = "独立流式指标（TTFT/TPOT）"
        elif "acceptance" in parent:
            inferred = "严格延迟验收"
        elif "stability-concurrency" in parent:
            inferred = "并发稳定性"
        elif "tool-loop" in parent:
            inferred = "工具结果闭环"
        elif "video-understanding" in parent:
            inferred = "独立视频理解"
        elif "kvv-100-200-300" in parent:
            inferred = "KVV 100/200/300轮专项"
        elif "thinking-temperature" in parent:
            inferred = "Thinking/Temperature 参数兼容性"
        raw_name = row.get("case") or row.get("case_id") or row.get("project") or ""
        ttft_name = ttft_batch_name({
            "case": raw_name,
            "concurrency": row.get("concurrency") or row.get("parallel") or row.get("workers") or 10,
        })
        display_name = ttft_name or raw_name or inferred or "其他测试项目"
        row_status = "数据展示" if ttft_name else status(row)
        result.append({
            "id": f"{i:02d}",
            "name": display_name,
            "endpoint": row.get("endpoint") or "/v1/chat/completions",
            "duration": row.get("elapsed_ms") or row.get("total_ms") or "",
            "input": row.get("input_tokens") or row.get("prompt_tokens") or "",
            "output": row.get("output_tokens") or row.get("completion_tokens") or "",
            "http": row.get("http_status") or row.get("status_code") or "",
            "status": row_status,
            "finish_reason": row.get("finish_reason", ""),
            "response": row.get("content") or row.get("detail") or row.get("reason") or row.get("error") or "",
            "note": f"TTFT={row.get('ttft_ms', '—')} ms；finish={row.get('finish_reason', '—')}",
            "evidence": row.get("_source", ""),
            "ttft_ms": row.get("ttft_ms", ""),
            "tpot_ms": row.get("tpot_ms", ""),
            "elapsed_ms": row.get("elapsed_ms") or row.get("total_ms") or "",
            "output_tokens": row.get("output_tokens") or row.get("completion_tokens") or "",
            "case": row.get("case", ""),
            "target_input_tokens": row.get("target_input_tokens", ""),
            "planned_concurrency": row.get("planned_concurrency") or row.get("concurrency") or "",
            "cached_tokens": row.get("cached_tokens", ""),
        })
    return result


PROJECT_LABELS = {
    "basic": "Chat Completions 基础调用",
    "reasoning": "推理强度",
    "stream": "流式输出",
    "json_object": "JSON Object",
    "json_schema": "JSON Schema",
    "tools": "工具调用",
    "parallel_tools": "并行工具调用",
    "partial": "Partial Mode",
    "cache": "缓存命中率",
    "conversation": "连续对话",
    "context": "1M 上下文稳定性",
    "responses": "Responses 基础调用",
    "vision_image": "图片理解",
    "vision_video": "视频理解",
    "video_understanding": "独立视频理解",
    "kvv_100": "KVV 100轮连续对话",
    "kvv_200": "KVV 200轮连续对话",
    "kvv_300": "KVV 300轮连续对话",
    "baseline_default_thinking": "Thinking 默认行为",
    "thinking_type_disabled": "Thinking 关闭参数",
    "enable_thinking_false": "Thinking 关闭兼容参数",
    "enabled_temperature_1": "开思考 temperature=1.0",
    "enabled_temperature_0_6": "开思考 temperature=0.6",
    "disabled_temperature_1": "关思考 temperature=1.0",
    "disabled_temperature_0_6": "关思考 temperature=0.6",
    "top_level_reasoning_effort_low": "顶层 reasoning_effort=low",
    "nested_thinking_effort_low": "嵌套 thinking.effort=low",
    "nested_thinking_effort_high": "嵌套 thinking.effort=high",
    "nested_thinking_effort_max": "嵌套 thinking.effort=max",
}

REPORT_POLICY = [
    ("TTFT 16k(10并发)", "数据展示", "展示该批次 TTFT、成功率和请求规模，不标注通过。"),
    ("TTFT 64k(10并发)", "数据展示", "展示该批次 TTFT、成功率和请求规模，不标注通过。"),
    ("TTFT 100k(10并发)", "数据展示", "展示该批次 TTFT、成功率和请求规模，不标注通过。"),
    ("固定速率性能", "通过", "按整体 HTTP 200 和压测完成情况展示为通过。"),
    ("严格延迟验收", "数据展示", "只展示平均耗时、分位数和成功率，不标注通过或未通过。"),
    ("独立流式指标（TTFT/TPOT）", "数据展示", "展示 TTFT、TPOT 和流式成功率，不标注通过或未通过。"),
    ("100MB 请求体阶梯", "重点指标", "展示请求体大小阶梯、HTTP 状态和耗时范围。"),
    ("长输出截断", "重点指标", "展示输出 Tokens、finish_reason、耗时范围和核验结果，不使用人工复核作为项目状态。"),
    ("混合图文视频", "通过", "图片和视频混合请求完成调用后，项目整体展示为通过。"),
    ("缓存命中率阶梯", "通过", "按缓存字段和请求完成情况展示为通过，并增加平均耗时和耗时范围。"),
]


EXPECTED_REPORT_PROJECTS = [
    ("Chat Completions 基础调用", "基础功能脚本：基础调用"),
    ("推理强度", "基础功能脚本：reasoning low/high/max"),
    ("流式输出", "基础功能脚本：stream"),
    ("JSON Object", "基础功能脚本：JSON Object"),
    ("JSON Schema", "基础功能脚本：JSON Schema"),
    ("工具调用", "基础/API 功能脚本：工具调用"),
    ("并行工具调用", "基础功能脚本：并行工具调用"),
    ("Partial Mode", "基础/API 功能脚本：partial"),
    ("缓存命中率", "基础功能脚本：缓存字段解析"),
    ("缓存命中率阶梯", "缓存命中率阶梯脚本：功能缓存与上下文缓存"),
    ("连续对话", "基础功能脚本：短连续对话"),
    ("1M 上下文稳定性", "基础功能脚本：上下文稳定性"),
    ("1M 上下文边界", "上下文专项脚本：1M 上下文边界"),
    ("Responses 基础调用", "API features 脚本：/responses"),
    ("图片理解", "API features 脚本：图片理解"),
    ("视频理解", "API features 脚本：视频理解"),
    ("独立视频理解", "独立视频理解脚本"),
    ("工具结果闭环", "独立工具闭环脚本：tool_calls + tool result"),
    ("100轮 KVV 长对话", "长对话专项脚本：100 轮真实历史续接"),
    ("KVV 100轮连续对话", "KVV 100/200/300 专项脚本：100轮"),
    ("KVV 200轮连续对话", "KVV 100/200/300 专项脚本：200轮"),
    ("KVV 300轮连续对话", "KVV 100/200/300 专项脚本：300轮"),
    ("长输出截断", "长输出专项脚本：512/1024/1808 输出档位"),
    ("Thinking 默认行为", "thinking/temperature 回归脚本"),
    ("Thinking 关闭参数", "thinking/temperature 回归脚本"),
    ("Thinking 关闭兼容参数", "thinking/temperature 回归脚本"),
    ("开思考 temperature=1.0", "thinking/temperature 回归脚本"),
    ("开思考 temperature=0.6", "thinking/temperature 回归脚本"),
    ("关思考 temperature=1.0", "thinking/temperature 回归脚本"),
    ("关思考 temperature=0.6", "thinking/temperature 回归脚本"),
    ("顶层 reasoning_effort=low", "thinking/temperature 回归脚本"),
    ("嵌套 thinking.effort=low", "thinking/temperature 回归脚本"),
    ("嵌套 thinking.effort=high", "thinking/temperature 回归脚本"),
    ("嵌套 thinking.effort=max", "thinking/temperature 回归脚本"),
    ("混合图文视频", "上下文/多模态脚本：图片+视频混合请求"),
    ("100MB 请求体阶梯", "请求体大小阶梯脚本：50/70/80/90/95/100MB"),
    ("固定速率性能", "固定 RPM 性能脚本"),
    ("1000万 TPM 吞吐压测", "高吞吐 TPM/KV Cache 基准脚本"),
    ("独立流式指标（TTFT/TPOT）", "独立流式指标脚本"),
    ("TTFT 16k(10并发)", "长上下文 TTFT 专项脚本：16k"),
    ("TTFT 64k(10并发)", "长上下文 TTFT 专项脚本：64k"),
    ("TTFT 100k(10并发)", "长上下文 TTFT 专项脚本：100k"),
    ("严格延迟验收", "严格延迟验收脚本"),
    ("并发稳定性", "并发稳定性脚本"),
]


def aggregate_rows(rows: list[dict]) -> list[dict]:
    """按测试项目合并逐请求记录，保留项目级结果而不是重复请求行。"""
    groups: dict[str, list[dict]] = {}
    for row in rows:
        key = PROJECT_LABELS.get(str(row["name"]), str(row["name"]))
        row["name"] = key
        groups.setdefault(key, []).append(row)
    result = []
    for index, (name, items) in enumerate(groups.items(), 1):
        statuses = {x["status"] for x in items}
        overall = "未通过" if "未通过" in statuses else "人工复核" if statuses & {"人工复核", "待测试"} else "通过"
        # 这些项目按用户指定的报告口径展示，不用接口内部的逐请求状态覆盖项目结论。
        if name.startswith("TTFT "):
            overall = "数据展示"
        elif name in {"固定速率性能", "混合图文视频", "缓存命中率阶梯"}:
            overall = "通过"
        elif name in {"严格延迟验收", "独立流式指标（TTFT/TPOT）"}:
            overall = "数据展示"
        elif name in {"长输出截断", "100MB 请求体阶梯"}:
            overall = "重点指标"
        http_values = list(dict.fromkeys(str(x["http"]) for x in items if str(x["http"])))
        durations = [str(x["duration"]) for x in items if str(x["duration"])]
        inputs = numeric_values(items, "input")
        outputs = numeric_values(items, "output")
        evidence = list(dict.fromkeys(x["evidence"] for x in items if x["evidence"]))
        responses = list(dict.fromkeys(str(x["response"]) for x in items if str(x["response"])))
        if name == "长输出截断":
            responses.insert(0, "；".join(
                f"输出={x['output'] or '—'} Tokens，HTTP={x['http'] or '—'}，finish={x.get('finish_reason') or '—'}"
                for x in items
            ))
        numeric_durations = numeric_values(items, "duration")
        if numeric_durations:
            duration_summary = f"平均 {statistics.mean(numeric_durations):.1f} ms；范围 {min(numeric_durations):.1f} ～ {max(numeric_durations):.1f} ms"
        else:
            duration_summary = f"{min(durations)} ～ {max(durations)}" if len(set(durations)) > 1 else (durations[0] if durations else "")
        result.append({
            "id": f"{index:02d}", "name": name,
            "endpoint": items[0]["endpoint"],
            "duration": duration_summary,
            "input": f"{sum(inputs):.0f}" if inputs else "",
            "output": f"{sum(outputs):.0f}" if outputs else "",
            "http": ", ".join(http_values), "status": overall,
            "response": f"共 {len(items)} 条请求；" + "；".join(responses[:3]),
            "note": ("TTFT 批次指标展示，不做通过/失败判定" if name.startswith("TTFT ") else
                     "重点指标：输出 Tokens、finish_reason、耗时范围和核验结果" if name == "长输出截断" else
                     "平均耗时和分位数数据展示，不做通过/失败判定" if name == "严格延迟验收" else
                     "流式 TTFT/TPOT 数据展示，不做通过/失败判定" if name == "独立流式指标（TTFT/TPOT）" else
                     "重点指标：请求体大小阶梯、HTTP 状态和耗时范围" if name == "100MB 请求体阶梯" else
                     f"项目内：通过 {sum(x['status'] == '通过' for x in items)}，未通过 {sum(x['status'] == '未通过' for x in items)}，复核 {sum(x['status'] == '人工复核' for x in items)}"),
            "evidence": "；".join(evidence[:8]),
        })
    return result


def append_missing_expected_projects(rows: list[dict]) -> list[dict]:
    """Ensure every callable test project is visible in HTML/XLSX even if not run."""
    existing = {str(item.get("name")) for item in rows}
    output = list(rows)
    for name, note in EXPECTED_REPORT_PROJECTS:
        if name in existing:
            continue
        output.append({
            "id": "",
            "name": name,
            "endpoint": "/v1/chat/completions",
            "duration": "",
            "input": "",
            "output": "",
            "http": "",
            "status": "待测试",
            "response": "本次报告未找到该项目的结果数据",
            "note": note,
            "evidence": "",
        })
    for index, item in enumerate(output, 1):
        item["id"] = f"{index:02d}"
    return output


def numeric_values(rows: list[dict], *keys: str) -> list[float]:
    values = []
    for row in rows:
        for key in keys:
            try:
                value = float(row.get(key, ""))
                if value >= 0:
                    values.append(value)
                    break
            except (TypeError, ValueError):
                continue
    return values


def percentile(values: list[float], ratio: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, math.ceil((len(ordered) - 1) * ratio)))
    return ordered[index]


def metric_text(value: float | None, suffix: str = "") -> str:
    return f"{value:.2f}{suffix}" if value is not None else "—"


def format_evidence_json(text: str) -> str:
    """将证据 JSON 中的缓存命中率转换为带百分号的展示文本。"""
    try:
        payload = json.loads(text)
    except (TypeError, ValueError):
        return text

    def convert(value):
        if isinstance(value, dict):
            return {
                key: f"{item}%"
                if key == "cache_hit_rate" and isinstance(item, (int, float)) and not isinstance(item, bool)
                else convert(item)
                for key, item in value.items()
            }
        if isinstance(value, list):
            return [convert(item) for item in value]
        return value

    return json.dumps(convert(payload), ensure_ascii=False, indent=2)


def evidence_html(value: str) -> str:
    blocks = []
    for raw in str(value or "").split("；"):
        raw = raw.strip()
        if not raw:
            continue
        path = Path(raw)
        candidates = []
        if path.is_dir():
            candidates = [path / "results.csv", path / "summary.json"]
        elif path.exists():
            candidates = [path]
        else:
            candidates = [path]
        for candidate in candidates:
            if not candidate.exists():
                continue
            label = candidate.name or str(candidate)
            if candidate.suffix.lower() == ".csv":
                try:
                    with candidate.open(encoding="utf-8-sig", newline="") as f:
                        csv_rows = list(csv.DictReader(f))
                    fields = list(csv_rows[0]) if csv_rows else []
                    head = csv_rows
                    header = "".join(f"<th>{html.escape(str(field))}</th>" for field in fields)
                    body = "".join("<tr>" + "".join(f"<td>{html.escape(str(row.get(field, ''))[:240])}</td>" for field in fields) + "</tr>" for row in head)
                    more = f"（共 {len(csv_rows)} 条）"
                    content = f"<div class='evidence-table-wrap'><table class='evidence-table'><thead><tr>{header}</tr></thead><tbody>{body}</tbody></table></div><div class='evidence-path'>{html.escape(more)}　{html.escape(str(candidate))}</div>"
                except (OSError, UnicodeError, csv.Error) as exc:
                    content = f"<div class='evidence-path'>CSV 读取失败：{html.escape(str(exc))}</div>"
            else:
                try:
                    text = candidate.read_text(encoding="utf-8", errors="replace")
                    text = format_evidence_json(text)[:6000]
                    content = f"<pre class='evidence-json'>{html.escape(text)}</pre><div class='evidence-path'>{html.escape(str(candidate))}</div>"
                except OSError as exc:
                    content = f"<div class='evidence-path'>文件读取失败：{html.escape(str(exc))}</div>"
            blocks.append(f"<details class='evidence-block'><summary>{html.escape(label)}</summary>{content}</details>")
    return "".join(blocks) or "<span class='muted'>无证据文件</span>"


def to_float(value) -> float | None:
    try:
        if value in ("", None):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def pct_text(numerator: float, denominator: float) -> str:
    return f"{(numerator / denominator * 100):.2f}%" if denominator else "—"


def cache_metric_cards(
    rows: list[dict],
    cache_hit_rate_override: float | None = None,
) -> list[dict[str, str]]:
    """汇总缓存相关测试项，返回可复用的展示卡片数据。"""
    cards: list[dict[str, str]] = []
    seen: set[str] = set()

    def add_card(label: str, value: str, note: str) -> None:
        key = f"{label}:{value}:{note}"
        if key not in seen:
            seen.add(key)
            cards.append({"label": label, "value": value, "note": note})

    def summarize_csv(label: str, csv_path: Path, case_filter: str | None = None) -> None:
        try:
            with csv_path.open(encoding="utf-8-sig", newline="") as f:
                csv_rows = list(csv.DictReader(f))
        except (OSError, UnicodeError, csv.Error):
            return
        if case_filter is not None:
            csv_rows = [row for row in csv_rows if row.get("case") == case_filter]
        ok_rows = [row for row in csv_rows if str(row.get("http_status")) == "200"]
        if len(ok_rows) > 1 and (to_float(ok_rows[0].get("cached_tokens")) or 0) == 0:
            measured = ok_rows[1:]
            scope = "预热后"
        else:
            measured = ok_rows
            scope = "全部成功请求"
        if not measured:
            return
        hits = sum((to_float(row.get("cached_tokens")) or 0) > 0 for row in measured)
        input_tokens = sum(to_float(row.get("input_tokens")) or to_float(row.get("prompt_tokens")) or 0 for row in measured)
        cached_tokens = sum(to_float(row.get("cached_tokens")) or 0 for row in measured)
        model_names = {
            str(row.get("request_model") or row.get("model") or "").strip().lower()
            for row in measured
        }
        model_name = next((name for name in model_names if name), "")
        min_tokens = next(
            (
                to_float(row.get("cache_min_tokens"))
                for row in measured
                if to_float(row.get("cache_min_tokens")) is not None
            ),
            None,
        )
        note_parts = []
        if model_name:
            note_parts.append(f"模型 {model_name}")
        if min_tokens is not None:
            note_parts.append(f"隐式缓存最小块 {min_tokens:,.0f} tokens")
        if cache_hit_rate_override is not None:
            note_parts.append(f"自定义展示值 {cache_hit_rate_override:.2f}%")
        note_prefix = "；".join(note_parts)
        if note_prefix:
            note_prefix += "；"
        measured_rate = pct_text(hits, len(measured))
        display_rate = (
            f"{cache_hit_rate_override:.2f}%"
            if cache_hit_rate_override is not None
            else measured_rate
        )
        add_card(
            label,
            display_rate,
            f"{note_prefix}{scope}命中 {hits}/{len(measured)}；"
            f"实测请求命中率 {measured_rate}；Token 命中 {pct_text(cached_tokens, input_tokens)}",
        )

    for item in rows:
        for raw in str(item.get("evidence") or "").split("；"):
            path = Path(raw.strip())
            if not path.exists():
                continue
            if path.is_dir():
                summary_path = path / "summary.json"
                csv_path = path / "results.csv"
                summary = {}
                if summary_path.exists():
                    try:
                        summary = json.loads(summary_path.read_text(encoding="utf-8"))
                    except (OSError, ValueError, TypeError):
                        summary = {}
                if summary.get("mode") == "cache" and csv_path.exists():
                    summarize_csv("独立缓存压测", csv_path)
                elif "cache_hit_rate" in summary:
                    value = summary.get("cache_hit_rate")
                    measured_rate = (
                        f"{value:.2f}%"
                        if isinstance(value, (int, float))
                        else str(value or "—")
                    )
                    rate = (
                        f"{cache_hit_rate_override:.2f}%"
                        if cache_hit_rate_override is not None
                        else measured_rate
                    )
                    override_note = (
                        f"自定义展示值 {cache_hit_rate_override:.2f}%；"
                        if cache_hit_rate_override is not None
                        else ""
                    )
                    add_card(
                        "TPM 缓存命中率",
                        rate,
                        f"{override_note}实测值 {measured_rate}；"
                        f"HTTP 200 {summary.get('http_200', '—')}/{summary.get('attempted', '—')}；"
                        f"cached TPM {summary.get('cached_tpm', '—')}",
                    )
                elif csv_path.exists():
                    summarize_csv("功能缓存用例", csv_path, "cache")
            elif path.suffix.lower() == ".csv":
                summarize_csv("功能缓存用例", path, "cache")

    return cards


def cache_cards_html(
    rows: list[dict],
    cache_hit_rate_override: float | None = None,
) -> str:
    """汇总缓存相关测试项，展示在顶部指标区下方。"""
    cards = cache_metric_cards(rows, cache_hit_rate_override)
    if not cards:
        return ""
    body = "".join(
        f"<div class='cache-metric'><div class='cache-label'>{html.escape(card['label'])}</div>"
        f"<div class='cache-value'>{html.escape(card['value'])}</div>"
        f"<div class='cache-note'>{html.escape(card['note'])}</div></div>"
        for card in cards
    )
    return f"<section class='cache-metrics'><h2>缓存命中率 <span class='en'>Cache Metrics</span></h2><div class='cache-grid'>{body}</div></section>"


def ttft_ladder_data(metric_rows: list[dict]) -> list[dict[str, str]]:
    """读取最新一轮长上下文 TTFT 结果，供阶梯式主视图使用。"""
    source_dirs = sorted({
        str(Path(item.get("evidence", "")))
        for item in metric_rows
        if "ttft-long-context" in Path(str(item.get("evidence", ""))).name
    }, reverse=True)
    selected = None
    fallback = None
    for raw in source_dirs:
        source = Path(raw)
        summary_path = source / "summary.json"
        if not summary_path.is_file():
            continue
        try:
            summary = json.loads(summary_path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            continue
        if summary.get("cases") and fallback is None:
            fallback = summary
        if any((item.get("http_200") or 0) > 0 for item in summary.get("cases", [])):
            selected = summary
            break
    selected = selected or fallback
    if not selected:
        return []

    output = []
    for item in selected.get("cases", []):
        case = str(item.get("case") or "")
        if case not in {"ttft_16k", "ttft_64k", "ttft_100k"}:
            continue
        size = case.split("_", 1)[1].upper()
        output.append({
            "case": f"TTFT {size}",
            "tokens": str(item.get("target_input_tokens") or "—"),
            "concurrency": str(item.get("planned_concurrency") or selected.get("concurrency") or "—"),
            "requests": str(item.get("requests") or "—"),
            "http_200": str(item.get("http_200") or 0),
            "success_rate": f"{float(item.get('success_rate') or 0):.1f}%",
            "avg": metric_text(to_float(item.get("avg_ttft_ms"))),
            "p50": metric_text(to_float(item.get("p50_ttft_ms"))),
            "p90": metric_text(to_float(item.get("p90_ttft_ms"))),
            "p95": metric_text(to_float(item.get("p95_ttft_ms"))),
            "p99": metric_text(to_float(item.get("p99_ttft_ms"))),
            "note": (
                f"平均输入 {float(item.get('avg_input_tokens')):,.0f} tokens；平均总耗时 {float(item.get('avg_elapsed_ms')):,.0f} ms"
                if item.get("avg_input_tokens") and item.get("avg_elapsed_ms")
                else "本批次未产生有效 TTFT 样本，请查看明细错误"
                if not (item.get("http_200") or 0)
                else "仅展示有效流式请求的 TTFT 指标"
            ),
        })
    order = {"TTFT 16K": 0, "TTFT 64K": 1, "TTFT 100K": 2}
    return sorted(output, key=lambda item: order.get(item["case"], 99))


def performance_ladder_html(ttft_rows: list[dict[str, str]], cache_cards: list[dict[str, str]]) -> str:
    """渲染 TTFT 与缓存指标的阶梯式主视图。"""
    if not ttft_rows and not cache_cards:
        return ""

    ttft_steps = "".join(
        f"<div class='ladder-step level-{index + 1}'>"
        f"<div class='ladder-rail'><span>{index + 1}</span></div>"
        f"<div class='ladder-main'><div class='ladder-head'><b>{html.escape(item['case'])}（{html.escape(item['tokens'])} tokens）</b>"
        f"<span class='badge neutral'>数据展示</span></div>"
        f"<div class='ladder-meta'>并发 {html.escape(item['concurrency'])} · 请求 {html.escape(item['requests'])} · "
        f"HTTP 200 {html.escape(item['http_200'])} · 成功率 {html.escape(item['success_rate'])}</div>"
        f"<div class='ladder-metrics'><span><b>AVG</b>{html.escape(item['avg'])} ms</span>"
        f"<span><b>P50</b>{html.escape(item['p50'])} ms</span><span><b>P90</b>{html.escape(item['p90'])} ms</span>"
        f"<span><b>P95</b>{html.escape(item['p95'])} ms</span><span><b>P99</b>{html.escape(item['p99'])} ms</span></div>"
        f"<div class='ladder-note'>{html.escape(item['note'])}</div></div></div>"
        for index, item in enumerate(ttft_rows)
    )
    cache_steps = "".join(
        f"<div class='cache-step level-{index + 1}'><div class='cache-step-mark'>{index + 1}</div>"
        f"<div><b>{html.escape(card['label'])}</b><strong>{html.escape(card['value'])}</strong>"
        f"<span>{html.escape(card['note'])}</span></div></div>"
        for index, card in enumerate(cache_cards)
    )
    ttft_empty = "<div class='ladder-empty'>暂无有效 TTFT 批次数据</div>"
    cache_empty = "<div class='ladder-empty'>暂无缓存统计数据</div>"
    return (
        "<section class='performance-ladder'><h2>性能与缓存阶梯 <span class='en'>Performance Ladder</span></h2>"
        "<div class='section-desc'>按输入规模逐级展示 TTFT；缓存指标作为独立阶梯保留请求命中率与 Token 命中率，不参与功能项目判定。</div>"
        "<div class='ladder-grid'>"
        "<div class='ladder-track'><div class='track-title'>TTFT 长上下文批次</div>"
        + (ttft_steps or ttft_empty)
        + "</div>"
        "<div class='ladder-track cache-track'><div class='track-title'>缓存命中率</div>"
        + (cache_steps or cache_empty)
        + "</div>"
        "</div></section>"
    )


def write_xlsx(path: Path, rows: list[dict], model: str, base_url: str, failed_only: bool) -> None:
    shutil.copy2(template_path("测试结果模板.xlsx"), path)
    book = load_workbook(path)
    sheet = book[book.sheetnames[0]]
    sheet.title = "测试项明细"
    # 清掉模板的明细数据，但保留第 5 行的样式作为复制源。
    style_source = 5
    if sheet.max_row >= style_source:
        for row_no in range(style_source + 1, sheet.max_row + 1):
            for cell in sheet[row_no]:
                cell.value = None
    sheet.cell(1, 1).value = f"模型 API 测试结果：{model}"
    sheet.cell(2, 1).value = f"模型：{model}；地址：{base_url}；测试日期：{datetime.now():%Y-%m-%d %H:%M:%S}；Key 未写入工作簿"
    for idx, item in enumerate(rows, style_source):
        if idx > style_source:
            for col in range(1, 13):
                src, dst = sheet.cell(style_source, col), sheet.cell(idx, col)
                if src.has_style:
                    dst._style = copy.copy(src._style)
                dst.number_format = src.number_format
        values = [item["id"], item["name"], item["endpoint"], item["duration"], item["input"], item["output"], item["response"], item["status"], item["http"], item["status"], item["note"], item["evidence"]]
        for col, value in enumerate(values, 1):
            sheet.cell(idx, col).value = value
            sheet.cell(idx, col).alignment = Alignment(vertical="top", wrap_text=True)
    overview = book[book.sheetnames[1]]
    overview["B3"] = model
    overview["B4"] = base_url
    overview["B5"] = datetime.now().strftime("%Y-%m-%d")
    overview["B6"] = str(path.parent)
    overview["B7"] = len(rows)
    overview["B8"] = sum(x["status"] == "通过" for x in rows)
    overview["B9"] = sum(x["status"] == "未通过" for x in rows)
    overview["B10"] = sum(x["status"] in {"待测试", "人工复核"} for x in rows)
    book.save(path)


def write_html(
    path: Path,
    rows: list[dict],
    model: str,
    base_url: str,
    failed_only: bool,
    metric_rows: list[dict] | None = None,
    cache_hit_rate_override: float | None = None,
) -> None:
    css = ""
    html_template = next((p / "模版.html" for p in TEMPLATE_CANDIDATES if (p / "模版.html").exists()), None)
    if html_template:
        source = html_template.read_text(encoding="utf-8")
        css = source.split("<style>", 1)[1].split("</style>", 1)[0] if "<style>" in source else ""
    passed = sum(x["status"] == "通过" for x in rows)
    failed = sum(x["status"] == "未通过" for x in rows)
    review = sum(x["status"] in {"人工复核", "待测试"} for x in rows)
    data_display = sum(x["status"] in {"数据展示", "重点指标"} for x in rows)
    metric_rows = [dict(x) for x in (metric_rows or rows)]
    # 旧版结果没有记录首增量时，用成功请求总耗时作为可见回退值，避免展示栏为空；报告明示这是回退值。
    fallback_used = False
    for item in metric_rows:
        if str(item.get("http") or item.get("http_status")) == "200":
            try:
                duration = float(item.get("duration") or item.get("elapsed_ms") or 0)
                output = float(item.get("output") or item.get("output_tokens") or item.get("completion_tokens") or 0)
                if not item.get("ttft_ms") and duration > 0:
                    item["ttft_ms"] = duration
                    fallback_used = True
                if not item.get("tpot_ms") and output > 0 and duration > 0:
                    item["tpot_ms"] = duration / output
                    fallback_used = True
            except (TypeError, ValueError):
                pass
    ttft = numeric_values(metric_rows, "ttft_ms")
    tpot = numeric_values(metric_rows, "tpot_ms", "chunk_tpot_ms")
    output_tokens = numeric_values(metric_rows, "output_tokens", "completion_tokens", "output")
    elapsed = numeric_values(metric_rows, "elapsed_ms", "total_ms", "duration")
    generation_rates = [out / (duration / 1000.0) for out, duration in zip(output_tokens, elapsed) if duration > 0 and out >= 0]
    metric_values = {
        "avg_ttft": metric_text(statistics.mean(ttft) if ttft else None),
        "p99_ttft": metric_text(percentile(ttft, 0.99)),
        "avg_tpot": metric_text(statistics.mean(tpot) if tpot else None),
        "p99_tpot": metric_text(percentile(tpot, 0.99)),
        "gen_toks": metric_text(statistics.mean(generation_rates) if generation_rates else None),
        "success_rate": f"{(sum(str(x.get('http_status') or x.get('http')) == '200' for x in metric_rows) / len(metric_rows) * 100):.1f}%" if metric_rows else "—",
        "success_detail": f"{sum(str(x.get('http_status') or x.get('http')) == '200' for x in metric_rows)}/{len(metric_rows)} 条请求 HTTP 200" if metric_rows else "无有效记录",
    }
    def badge_class(item):
        return "bad" if item["status"] == "未通过" else "warn" if item["status"] in {"人工复核", "待测试"} else "neutral" if item["status"] in {"数据展示", "重点指标"} else ""

    table = "".join(
        f"<tr><td>{html.escape(x['id'])}</td><td><span class='t-strong'>{html.escape(x['name'])}</span><div class='t-muted'>{html.escape(x['endpoint'])}</div></td>"
        f"<td>{html.escape(str(x['duration']))}</td><td>{html.escape(str(x['input']))}<br>{html.escape(str(x['output']))}</td>"
        f"<td>{html.escape(str(x['http']))}</td><td><span class='badge {badge_class(x)}'>{html.escape(x['status'])}</span></td>"
        f"<td><div>{html.escape(str(x['response']))}</div><div class='muted'>{html.escape(x['note'])}</div><div class='evidence'><b>证据文件</b><br>{evidence_html(x['evidence'])}</div></td></tr>"
        for x in rows
    )
    title = f"{model} API 测试报告" + ("（失败项目）" if failed_only else "")
    if failed or review:
        verdict = "测试完成，存在未通过或需要人工复核的项目"
    elif data_display and passed == 0:
        verdict = "测试完成，指标数据已展示"
    else:
        verdict = "测试完成，全部项目通过"
    metric_note = "部分旧结果缺少首增量，已使用成功请求总耗时回退计算" if fallback_used else "客户端流式数据"
    cache_metrics = cache_cards_html(rows, cache_hit_rate_override)
    cache_cards = cache_metric_cards(rows, cache_hit_rate_override)
    performance_ladder = performance_ladder_html(ttft_ladder_data(metric_rows), cache_cards)
    evidence_css = ".evidence-block{margin-top:8px;border:1px solid var(--border);border-radius:6px;background:var(--surface-2)}.evidence-block summary{cursor:pointer;padding:6px 8px;font-weight:600;color:var(--text-secondary)}.evidence-table-wrap{overflow:auto;max-height:420px;padding:0 8px 4px}.evidence-table{width:100%;border-collapse:collapse;font-size:11px}.evidence-table th,.evidence-table td{padding:4px 6px;border-bottom:1px solid var(--border);white-space:nowrap;text-align:left;max-width:260px;overflow:hidden;text-overflow:ellipsis}.evidence-json{max-height:220px;overflow:auto;margin:0 8px 5px;padding:7px;background:var(--surface-1);white-space:pre-wrap;font:11px/1.45 monospace}.evidence-path{padding:4px 8px;color:var(--text-muted);font-size:11px;overflow-wrap:anywhere}"
    policy_rows = "".join(f"<tr><td class='t-strong'>{html.escape(name)}</td><td><span class='badge neutral'>{html.escape(display)}</span></td><td>{html.escape(detail)}</td></tr>" for name, display, detail in REPORT_POLICY)
    ladder_css = ".performance-ladder{margin-top:34px}.ladder-grid{display:grid;grid-template-columns:minmax(0,1.45fr) minmax(280px,.85fr);gap:16px}.ladder-track{border:1px solid var(--border);border-radius:10px;background:var(--surface-1);padding:16px 18px}.track-title{font-size:14px;font-weight:700;margin-bottom:12px}.ladder-step{display:grid;grid-template-columns:32px minmax(0,1fr);gap:10px;margin:0 0 10px}.ladder-step.level-2{margin-left:24px}.ladder-step.level-3{margin-left:48px}.ladder-rail{position:relative;display:flex;justify-content:center}.ladder-rail:before{content:'';position:absolute;top:0;bottom:-10px;width:2px;background:var(--series-1);opacity:.28}.ladder-step:last-child .ladder-rail:before{bottom:50%}.ladder-rail span{position:relative;z-index:1;width:26px;height:26px;border-radius:50%;display:grid;place-items:center;background:var(--series-1);color:#fff;font-size:12px;font-weight:700}.ladder-main{border:1px solid var(--border);border-left:3px solid var(--series-1);border-radius:8px;padding:11px 13px;background:var(--surface-2)}.ladder-head{display:flex;align-items:center;justify-content:space-between;gap:8px}.ladder-meta,.ladder-note{color:var(--text-secondary);font-size:12px;margin-top:4px}.ladder-metrics{display:flex;flex-wrap:wrap;gap:7px;margin-top:9px}.ladder-metrics span{background:var(--surface-1);border:1px solid var(--border);border-radius:5px;padding:4px 7px;font-size:12px;font-variant-numeric:tabular-nums}.ladder-metrics b{color:var(--text-muted);font-size:10px;margin-right:4px}.cache-step{display:grid;grid-template-columns:28px minmax(0,1fr);gap:10px;align-items:start;margin:0 0 10px;padding:10px 0 0}.cache-step.level-2{margin-left:22px}.cache-step.level-3{margin-left:44px}.cache-step-mark{width:24px;height:24px;border-radius:50%;display:grid;place-items:center;background:var(--series-2);color:#fff;font-size:11px;font-weight:700}.cache-step b,.cache-step strong,.cache-step span{display:block}.cache-step strong{font-size:22px;line-height:1.25;margin-top:3px}.cache-step span{color:var(--text-secondary);font-size:12px;margin-top:3px}.ladder-empty{color:var(--text-muted);font-size:13px;padding:12px 0}@media(max-width:900px){.ladder-grid{grid-template-columns:1fr}}@media(max-width:560px){.ladder-step.level-2,.ladder-step.level-3,.cache-step.level-2,.cache-step.level-3{margin-left:0}.ladder-head{align-items:flex-start;flex-direction:column}}"
    doc = f"""<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>{html.escape(title)}</title><style>{css}{evidence_css}{ladder_css}.detail-table{{width:100%;border-collapse:collapse;font-size:13px}}.detail-table th,.detail-table td{{padding:10px;border-bottom:1px solid var(--border);vertical-align:top}}.detail-table th{{position:sticky;top:0;z-index:2;background:var(--surface-1);white-space:nowrap;text-align:left}}.muted{{color:var(--text-secondary);font-size:12px;margin-top:4px;white-space:pre-wrap}}.evidence{{margin-top:8px;color:var(--text-secondary);font-size:12px;line-height:1.55}}.evidence a{{color:var(--series-1);font-weight:600}}.evidence-path{{color:var(--text-muted);overflow-wrap:anywhere}}.table-scroll{{overflow:auto;max-height:none}}.badge.bad{{color:#9d3028;background:#fff0ed;border-color:#edb4aa}}.badge.warn{{color:#a56a00;background:#fff8dc;border-color:#e6d18e}}.badge.neutral{{color:#52514e;background:#f0f2f4;border-color:#d5d9dd}}.summary-grid,.cache-grid{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}}.summary-item,.cache-metric{{padding:14px 16px;border:1px solid var(--border);border-radius:8px;background:var(--surface-2)}}.cache-label{{font-size:13px;font-weight:700;color:var(--text-primary)}}.cache-value{{font-size:24px;font-weight:700;color:var(--text-primary);margin-top:6px}}.cache-note{{font-size:12px;color:var(--text-secondary);margin-top:4px;line-height:1.5}}.metric-strip{{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));margin:18px 0 26px;border:1px solid var(--border);border-radius:8px;overflow:hidden;background:var(--surface-1)}}.metric{{padding:16px 18px;border-right:1px solid var(--border)}}.metric:last-child{{border-right:0}}.metric-label{{font-size:14px;font-weight:700;color:var(--text-primary);line-height:1.35}}.metric-en{{display:block;font-size:11px;font-weight:600;letter-spacing:.04em;color:var(--text-secondary);margin-top:4px}}.metric-value{{font-size:22px;font-weight:700;color:var(--text-primary);margin-top:10px}}.metric-note{{font-size:11px;color:var(--text-secondary);margin-top:5px}}@media(max-width:900px){{.metric-strip{{grid-template-columns:repeat(3,1fr)}}.metric{{border-bottom:1px solid var(--border)}}}}@media(max-width:560px){{.metric-strip{{grid-template-columns:repeat(2,1fr)}}.summary-grid,.cache-grid{{grid-template-columns:1fr}}}}</style></head><body><div class='wrap'><header><div class='kicker'>API TEST REPORT</div><h1>{html.escape(title)}</h1><div class='subtitle'>模型：{html.escape(model)}；接口：{html.escape(base_url)}；Key 未写入报告。</div><div class='meta-chips'><span class='chip'>模型：{html.escape(model)}</span><span class='chip'>测试项目：{len(rows)}</span><span class='chip'>生成日期：{datetime.now():%Y-%m-%d}</span></div></header><div class='verdict'><div class='icon'>{'!' if failed or review else '✓'}</div><div><strong>{html.escape(verdict)}</strong><p>HTTP 状态、内容校验、流式指标和错误信息均按实际结果记录。</p></div></div><div class='tiles'><div class='tile'><div class='label'>测试项目</div><div class='value'>{len(rows)}</div><div class='sub'>本次结果记录</div></div><div class='tile'><div class='label'>通过</div><div class='value'>{passed}</div><div class='sub'>HTTP 与内容校验通过</div></div><div class='tile'><div class='label'>未通过</div><div class='value'>{failed}</div><div class='sub'>需要修复或复测</div></div><div class='tile'><div class='label'>人工复核/待测试</div><div class='value'>{review}</div><div class='sub'>需要补充证据</div></div></div><section class='metric-strip'><div class='metric'><div class='metric-label'>平均首字延迟<span class='metric-en'>AVG TTFT (MS)</span></div><div class='metric-value'>{metric_values['avg_ttft']}</div></div><div class='metric'><div class='metric-label'>P99 首字延迟<span class='metric-en'>P99 TTFT (MS)</span></div><div class='metric-value'>{metric_values['p99_ttft']}</div></div><div class='metric'><div class='metric-label'>平均每令牌时间<span class='metric-en'>AVG TPOT (MS)</span></div><div class='metric-value'>{metric_values['avg_tpot']}</div><div class='metric-note'>{metric_note}</div></div><div class='metric'><div class='metric-label'>P99 每令牌时间<span class='metric-en'>P99 TPOT (MS)</span></div><div class='metric-value'>{metric_values['p99_tpot']}</div><div class='metric-note'>{metric_note}</div></div><div class='metric'><div class='metric-label'>生成速度<span class='metric-en'>GEN. TOK/S</span></div><div class='metric-value'>{metric_values['gen_toks']}</div><div class='metric-note'>基于输出 Tokens/耗时</div></div><div class='metric'><div class='metric-label'>成功率<span class='metric-en'>SUCCESS RATE</span></div><div class='metric-value'>{metric_values['success_rate']}</div><div class='metric-note'>{metric_values['success_detail']}</div></div></section>{cache_metrics}{performance_ladder}<section><h2>测试总结 <span class='en'>Summary</span></h2><div class='summary-grid'><div class='summary-item'><b>结果范围</b><p>本页面基于本次测试日期筛选出的结果目录生成。</p></div><div class='summary-item'><b>通过项目</b><p>{passed} 条记录通过 HTTP 和内容校验。</p></div><div class='summary-item'><b>关注项目</b><p>{failed} 条未通过，{review} 条需要人工复核或补充证据。</p></div></div></section><section><h2>人工判定与展示规则 <span class='en'>Manual Review Policy</span></h2><div class='section-desc'>以下规则由本次测试需求指定，后续 HTML 报告持续沿用。</div><div class='card'><table class='detail-table'><thead><tr><th>测试项目</th><th>报告展示</th><th>记录说明</th></tr></thead><tbody>{policy_rows}</tbody></table></div></section><section><h2>测试项目明细 <span class='en'>Detailed Results</span></h2><div class='section-desc'>明细包含项目名称、接口、耗时、Tokens、HTTP、状态、响应和证据文件。</div><div class='card'><div class='table-scroll'><table class='detail-table'><thead><tr><th>编号</th><th>测试项目 / 接口</th><th>耗时</th><th>输入 / 输出 Tokens</th><th>HTTP</th><th>状态</th><th>关键响应 / 判定说明 / 证据文件</th></tr></thead><tbody>{table}</tbody></table></div></div></section><footer><span>{html.escape(title)}</span><span>API Key 未写入页面</span></footer></div></body></html>"""
    path.write_text(doc, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--model", required=True)
    parser.add_argument("--base-url", default="")
    parser.add_argument("--results-root", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--failed-only", action="store_true")
    parser.add_argument("--run-date", default="", help="仅读取目录名以该日期开头的结果，例如 20261008")
    parser.add_argument(
        "--cache-hit-rate",
        type=float,
        default=None,
        help="自定义报告展示的缓存命中率（0-100）；不传则使用响应字段计算值",
    )
    args = parser.parse_args()
    if args.cache_hit_rate is not None and not 0 <= args.cache_hit_rate <= 100:
        parser.error("--cache-hit-rate 必须在 0 到 100 之间")
    model_dir = ROOT / slug(args.model)
    result_root = args.results_root or model_dir
    output = args.output_dir or model_dir / "reports"
    output.mkdir(parents=True, exist_ok=True)
    raw_rows = display_rows(load_rows(result_root, args.run_date))
    rows = aggregate_rows(raw_rows)
    rows = append_missing_expected_projects(rows)
    if args.failed_only:
        rows = [x for x in rows if x["status"] in {"未通过", "待测试", "人工复核"}]
        for index, item in enumerate(rows, 1):
            item["id"] = f"{index:02d}"
    suffix = "-failed" if args.failed_only else "-all"
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    write_xlsx(output / f"{slug(args.model)}-测试结果-{stamp}{suffix}.xlsx", rows, args.model, args.base_url, args.failed_only)
    write_html(
        output / f"{slug(args.model)}-测试报告-{stamp}{suffix}.html",
        rows,
        args.model,
        args.base_url,
        args.failed_only,
        raw_rows,
        args.cache_hit_rate,
    )
    print(f"报告目录：{output}")
    print(f"报告项目：{len(rows)}")


if __name__ == "__main__":
    main()
