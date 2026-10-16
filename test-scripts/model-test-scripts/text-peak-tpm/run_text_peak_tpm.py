#!/usr/bin/env python3
"""纯文本峰值 TPM / 首字延迟实测，不设性能门槛。默认预览，不发送请求。"""
import argparse
import asyncio
import csv
import hashlib
import json
import math
import os
import platform
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
CASE = 'text_peak_tpm'
DEFAULTS = dict(base_url='http://127.0.0.1:8000/v1', key_env='MODEL_TEST_API_KEY',
                model='kimi-k3', duration=120., rate_rps=128., concurrency=640,
                input_units=32768, max_output_tokens=128, timeout=60., max_requests=16000,
                tpm_basis='total', cache_mode='shared', min_success_rate=.9)
ENV = dict(base_url='API_BASE', key_env='API_KEY_ENV', model='MODEL')
PURPOSE = '验证纯文本服务在负载下的真实峰值 TPM 和首字延迟，不发送视频或其他多模态内容'


def now():
    return datetime.now(ZoneInfo('Asia/Shanghai')).isoformat()


def dump(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def redact(obj, key):
    if isinstance(obj, str): return obj.replace(key, '[REDACTED]')
    if isinstance(obj, list): return [redact(v, key) for v in obj]
    if isinstance(obj, dict): return {k: redact(v, key) for k, v in obj.items()}
    return obj


def options(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--config', type=Path)
    for k in DEFAULTS:
        names = ['--' + k.replace('_', '-')]
        if k == 'model': names.append('--models')
        p.add_argument(*names, dest=k)
    p.add_argument('--cases', default=CASE)
    p.add_argument('--list', action='store_true')
    p.add_argument('--execute', '--run', action='store_true')
    p.add_argument('--confirm-paid', action='store_true')
    p.add_argument('--output-dir', type=Path)
    a = p.parse_args(argv)
    cfg = json.loads(a.config.read_text(encoding='utf-8')) if a.config else {}
    legacy = {'target_tpm', 'ttft_ms', 'ttft_percentile'}
    if not isinstance(cfg, dict) or set(cfg) - set(DEFAULTS) - legacy:
        raise ValueError('配置必须为对象，字段见 config.template.json')
    if set(cfg) & legacy:
        print('旧配置中的峰值/首字门槛已忽略，改为仅报告实测值。')
    for k, default in DEFAULTS.items():
        v = getattr(a, k)
        if v is None: v = os.getenv(ENV.get(k, 'TEXT_PEAK_' + k.upper()))
        if v is None: v = cfg.get(k, default)
        if isinstance(default, (float, int)):
            if isinstance(v, bool): raise ValueError(k + ' 不能是布尔值')
            if isinstance(default, int) and isinstance(v, float): raise ValueError(k + ' 必须为整数')
            v = type(default)(v)
            if not math.isfinite(v): raise ValueError(k + ' 必须为有限数')
        elif not isinstance(v, str) or not v.strip(): raise ValueError(k + ' 必须为非空字符串')
        setattr(a, k, v)
    limits = dict(duration=(60, 1800), rate_rps=(.01, 2000), concurrency=(1, 4096),
                  input_units=(128, 131072), max_output_tokens=(16, 8192), timeout=(6, 300),
                  max_requests=(2, 200000), min_success_rate=(.01, 1))
    for k, (lo, hi) in limits.items():
        if not lo <= getattr(a, k) <= hi: raise ValueError(f'{k} 范围为 {lo}–{hi}')
    if a.tpm_basis not in ('total', 'input', 'output', 'uncached_total'):
        raise ValueError('tpm-basis 仅支持 total/input/output/uncached_total')
    if a.cache_mode not in ('shared', 'unique'): raise ValueError('cache-mode 仅支持 shared/unique')
    u = urlsplit(a.base_url)
    if (u.scheme not in ('http', 'https') or not u.hostname or u.username or u.password
            or u.query or u.fragment): raise ValueError('base-url 必须为无鉴权信息及查询参数的 HTTP(S) 基础地址')
    _ = u.port
    if ',' in a.model or a.cases != CASE: raise ValueError('仅支持单模型和 text_peak_tpm 测试项')
    return a


def body(a, index, prefix):
    salt = f'Unique reference {index}.\n' if a.cache_mode == 'unique' else ''
    return dict(model=a.model, messages=[
        dict(role='system', content=salt + 'Read the reference; follow the final instruction.\n' + prefix),
        dict(role='user', content='Reply exactly PERF_OK and nothing else.')],
        max_tokens=a.max_output_tokens, stream=True, stream_options=dict(include_usage=True))


def new_record(index, scheduled):
    return dict(index=index, scheduled_s=scheduled, started_s=None, completed_s=None,
                http_status=None, success=False, ttft_ms=None, elapsed_ms=None,
                input_tokens=None, output_tokens=None, total_tokens=None, cached_tokens=None,
                finish_reason=None, sse_done=False, error_type=None, content='', usage={}, events=[])


def usage_counts(r):
    u = r['usage']
    def count(v): return v if type(v) is int and v >= 0 else None
    r['input_tokens'] = count(u.get('prompt_tokens'))
    r['output_tokens'] = count(u.get('completion_tokens'))
    # total 由两个官方计数相加；不使用提示词长度估算。
    if r['input_tokens'] is not None and r['output_tokens'] is not None:
        r['total_tokens'] = r['input_tokens'] + r['output_tokens']
    details = u.get('prompt_tokens_details') or {}
    for v in [details.get('cached_tokens'), u.get('prompt_cache_hit_tokens'), u.get('cached_tokens')]:
        if count(v) is not None:
            r['cached_tokens'] = v
            break


def parse_event(data, r, received_s):
    if data.strip() == '[DONE]':
        r['sse_done'] = True
        return
    item = json.loads(data)
    r['events'].append(item)
    if item.get('error'): raise ValueError('stream_error')
    if item.get('usage'):
        r['usage'] = item['usage']
        usage_counts(r)
    r['model_returned'] = item.get('model') or r.get('model_returned')
    for c in item.get('choices', []):
        if c.get('index', 0) != 0: continue
        piece = (c.get('delta') or {}).get('content')
        if isinstance(piece, str) and piece:
            if r['ttft_ms'] is None:
                r['ttft_ms'] = (received_s - r['started_s']) * 1000
            r['content'] += piece
        r['finish_reason'] = c.get('finish_reason') or r['finish_reason']


async def request(session, a, key, index, scheduled, origin, prefix, shared_payload):
    r = new_record(index, scheduled)
    payload = shared_payload if a.cache_mode == 'shared' else json.dumps(body(a, index, prefix)).encode('utf-8')
    r['started_s'] = time.perf_counter() - origin
    try:
        async with session.post(a.base_url.rstrip('/') + '/chat/completions', data=payload,
                                allow_redirects=False) as response:
            r.update(http_status=response.status, content_type=response.headers.get('Content-Type'),
                     request_id=response.headers.get('x-request-id'), decoding='utf-8',
                     ttfb_ms=(time.perf_counter() - origin - r['started_s']) * 1000)
            if response.status != 200:
                r['error_response'] = (await response.content.read(8192)).decode('utf-8', errors='replace')
                r['error_type'] = 'HTTPError'
            elif 'text/event-stream' not in (r['content_type'] or '').lower():
                r['error_type'] = 'NonSSEResponse'
            else:
                pending = []
                async for raw in response.content:
                    line = raw.decode('utf-8').rstrip('\r\n')
                    if not line:
                        if pending:
                            parse_event('\n'.join(pending), r, time.perf_counter() - origin)
                            pending = []
                        if r['sse_done']: break
                    elif line.startswith('data:'):
                        pending.append(line[5:].removeprefix(' '))
                if pending: parse_event('\n'.join(pending), r, time.perf_counter() - origin)
                r['success'] = (r['sse_done'] and r['finish_reason'] == 'stop'
                                and r['content'].strip() == 'PERF_OK')
    except Exception as exc:
        r['error_type'] = type(exc).__name__
        r['success'] = False
    r['completed_s'] = time.perf_counter() - origin
    r['elapsed_ms'] = (r['completed_s'] - r['started_s']) * 1000
    return redact(r, key)


def pct(values, p):
    if not values: return None
    values = sorted(values)
    # 最近秩百分位，P100 为最大值。
    return values[max(0, math.ceil(len(values) * p / 100) - 1)]


def metric(r, basis):
    if not r['success']: return None
    if basis == 'uncached_total':
        i, o, c = r['input_tokens'], r['output_tokens'], r['cached_tokens']
        if None in (i, o, c) or c > i: return None
        return i + o - c
    return r[{'total': 'total_tokens', 'input': 'input_tokens', 'output': 'output_tokens'}[basis]]


def peak_window(rows, basis, observed_s):
    """按完成时间计入完整 60 秒窗口，禁止短窗口外推成 TPM。"""
    if observed_s < 60: return dict(tokens=0, start_s=None, end_s=None)
    events = sorted((r['completed_s'], metric(r, basis)) for r in rows
                    if metric(r, basis) is not None and r['completed_s'] <= observed_s)
    left, total = 0, 0
    best = dict(tokens=0, start_s=0., end_s=60.)
    for right, (t, value) in enumerate(events):
        total += value
        end = max(60., t)
        while left <= right and events[left][0] <= end - 60:
            total -= events[left][1]
            left += 1
        if total > best['tokens']:
            best = dict(tokens=total, start_s=end - 60, end_s=end)
    return best


def summarize(rows, a, observed_s):
    peak = peak_window(rows, a.tpm_basis, observed_s)
    success = [r for r in rows if r['success']]
    values = [r['ttft_ms'] for r in success if r['ttft_ms'] is not None]
    coverage = all(metric(r, a.tpm_basis) is not None for r in success) and bool(success)
    success_rate = len(success) / len(rows) if rows else 0
    complete = observed_s >= 60 and coverage and len(values) == len(success)
    cohort = [r for r in rows if r['completed_s'] is not None
              and peak['start_s'] is not None and peak['start_s'] < r['completed_s'] <= peak['end_s']]
    return dict(measurement_completed=complete,
                success_rate_passed=success_rate >= a.min_success_rate,
                min_success_rate=a.min_success_rate,
                performance_mode='measurement_only',
                tpm_basis=a.tpm_basis, peak_60s=peak,
                peak_window_successes=sum(r['success'] for r in cohort),
                peak_window_failures=sum(not r['success'] for r in cohort),
                peak_by_basis={b: peak_window(rows, b, observed_s) for b in
                               ('input', 'output', 'total', 'uncached_total')},
                observed_s=observed_s, scheduled_requests=len(rows), sent_requests=sum(r['started_s'] is not None for r in rows),
                success_requests=len(success), success_rate=success_rate,
                client_dropped=sum(r['error_type'] in ('ClientCapacity', 'ClientScheduleLate') for r in rows),
                usage_coverage=len([r for r in success if metric(r, a.tpm_basis) is not None])/len(success) if success else 0,
                ttft_ms={f'P{p}': pct(values, p) for p in (50, 90, 95, 99, 100)},
                p99_ttft_ms=pct(values, 99),
                definitions=dict(tpm='成功请求官方 usage，按完成时间归属 (start,end] 的完整滚动60秒窗口；不外推',
                                 ttft='从开始发送至首个非空正文 delta.content，含连接及连接池等待',
                                 success='HTTP200 + SSE DONE + finish_reason=stop + 正文PERF_OK',
                                 cache='shared 使用相同前缀，可能命中缓存；unique 在前缀开头添加请求编号'))


async def run(a, key, out):
    import aiohttp
    prefix = 'x ' * a.input_units
    shared = json.dumps(body(a, 0, prefix)).encode('utf-8')
    origin = time.perf_counter()
    timeout = aiohttp.ClientTimeout(total=a.timeout)
    connector = aiohttp.TCPConnector(limit=a.concurrency, limit_per_host=a.concurrency)
    headers = {'Authorization': 'Bearer ' + key, 'Content-Type': 'application/json', 'Accept': 'text/event-stream'}
    async with aiohttp.ClientSession(connector=connector, timeout=timeout, headers=headers) as session:
        preflight = await request(session, a, key, 0, 0, origin, prefix, shared)
        dump(out / 'evidence' / 'preflight.json', preflight)
        # 大负载之前确认权限、额度、SSE及usage；探测失败不继续请求。
        if not preflight['success'] or metric(preflight, a.tpm_basis) is None:
            return [], 0., preflight
        origin = time.perf_counter()
        rows, active, tasks = [], set(), []
        count = math.ceil(a.duration * a.rate_rps)
        async def one(index, due):
            r = await request(session, a, key, index, due, origin, prefix, shared)
            # 单请求证据保留事件和请求生成规则；长前缀保存一次，避免海量重复文件。
            r.update(request_fixture='../request_template.json', unique_index=index,
                     case_id=CASE, purpose_zh=PURPOSE)
            dump(out / 'evidence' / f'{index:06}.json', r)
            rows.append(r)
        for i in range(count):
            due = i / a.rate_rps
            await asyncio.sleep(max(0., origin + due - time.perf_counter()))
            elapsed = time.perf_counter() - origin
            error = ('ClientScheduleLate' if elapsed - due > max(.05, 2/a.rate_rps)
                     else 'ClientCapacity' if len(active) >= a.concurrency else None)
            if error:
                r = new_record(i+1, due)
                r.update(error_type=error, completed_s=elapsed)
                rows.append(r)
            else:
                task = asyncio.create_task(one(i+1, due))
                tasks.append(task)
                active.add(task)
                task.add_done_callback(active.discard)
            if i and i % max(1, int(a.rate_rps * 5)) == 0:
                print(f'已计划={i}/{count}；在途={len(active)}；完成={len(rows)}', flush=True)
        await asyncio.sleep(max(0., origin + a.duration - time.perf_counter()))
        if tasks: await asyncio.gather(*tasks)
        return sorted(rows, key=lambda r: r['index']), time.perf_counter() - origin, preflight


def report(out, rows, summary, a):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    fields = sorted(set().union(*(r.keys() for r in rows)) - {'events', 'usage', 'content', 'error_response'}) if rows else ['index']
    with (out / 'results.csv').open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = '纯文本性能测试项'
    ws.merge_cells('A1:J1')
    ws['A1'] = f'{a.model} 峰值 TPM / 首字延迟实测：' + ('数据完整' if summary['measurement_completed'] else '数据不完整')
    ws.merge_cells('A2:J2')
    ws['A2'] = f'{a.base_url}；Key 已脱敏；{now()[:10]}'
    heads = ['编号','测试项目','接口','总耗时（秒）','Input Tokens','Output Tokens','关键响应','状态','主要作用','当前结果影响']
    for j,h in enumerate(heads,1): ws.cell(4,j,h)
    for i,r in enumerate(rows,5):
        values = [f'{r["index"]:02}', '纯文本负载请求', a.base_url.rstrip('/')+'/chat/completions',
                  r['elapsed_ms']/1000 if r['elapsed_ms'] is not None else None,
                  r['input_tokens'],r['output_tokens'],
                  f"HTTP={r['http_status']}；TTFT={r['ttft_ms']}ms；{r['error_type'] or r['finish_reason']}",
                  '成功' if r['success'] else '失败', PURPOSE,
                  '实际性能数据见实测汇总' if r['success'] else '此样本降低成功率，不能提供有效吞吐']
        for j,v in enumerate(values,1): ws.cell(i,j,v)
        ws.cell(i,4).number_format='0.000'
        ws.row_dimensions[i].height=48
    edge = Side(style='thin', color='B7C9D6')
    for row in ws:
        for c in row:
            c.font=Font(name='宋体',size=14 if c.row==1 else 10 if c.row==2 else 11,
                        bold=c.row in (1,4),color='FFFFFF' if c.row==4 else '000000')
            c.alignment=Alignment(vertical='center',wrap_text=True)
            if c.row>=4:
                c.border=Border(left=edge,right=edge,top=edge,bottom=edge)
                c.fill=PatternFill('solid',fgColor='4472C4' if c.row==4 else 'EAF1FB' if c.row%2 else 'FFFFFF')
    for j,width in enumerate([9,26,29,20,17,18,48,15,32,40],1):
        ws.column_dimensions[get_column_letter(j)].width=width
    for row,height in [(1,28),(2,24),(4,32)]: ws.row_dimensions[row].height=height
    ws.freeze_panes='A5'
    ws.auto_filter.ref=f'A4:J{max(4,ws.max_row)}'
    ws.print_area=f'A1:J{max(4,ws.max_row)}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth,ws.page_setup.fitToHeight=1,0
    overview=wb.create_sheet('实测汇总')
    for k,v in summary.items(): overview.append([k,json.dumps(v,ensure_ascii=False)])
    overview.column_dimensions['A'].width=35
    overview.column_dimensions['B'].width=100
    wb.save(out/'results.xlsx')


def main(argv=None):
    try: a=options(argv)
    except (ValueError, OSError, TypeError) as exc:
        print(f'参数错误：{exc}',file=sys.stderr)
        return 2
    if a.list:
        print(CASE+'：'+PURPOSE)
        return 0
    planned=math.ceil(a.duration*a.rate_rps)+1
    print(f'纯文本，不测视频；模型={a.model}；端点={a.base_url}/chat/completions')
    print(f'实测模式：峰值 TPM（{a.tpm_basis}）、P99 首字延迟不设门槛；成功率要求 ≥{a.min_success_rate:.0%}')
    print(f'负载={a.rate_rps:g} RPS；发起窗口={a.duration:g}s；并发上限={a.concurrency}；含探测计划={planned}；上限={a.max_requests}')
    print(f'输入单位={a.input_units} 个 x；长度仅为负载配置，TPM只用官方usage；缓存模式={a.cache_mode}')
    if planned>a.max_requests:
        print('计划超过 max-requests',file=sys.stderr)
        return 3
    if not a.execute and not a.confirm_paid:
        print('仅预览；实际执行需 --run --confirm-paid')
        return 0
    if not (a.execute and a.confirm_paid):
        print('实际执行需同时提供 --run --confirm-paid',file=sys.stderr)
        return 3
    key=os.environ.get(a.key_env)
    if not key or '\r' in key or '\n' in key:
        print('Key 环境变量未配置或格式无效；--key-env 必须传变量名',file=sys.stderr)
        return 2
    try:
        import aiohttp, openpyxl
    except ImportError:
        print('请先 python3 -m pip install -r requirements.txt',file=sys.stderr)
        return 2
    out=a.output_dir or ROOT/'results'/datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d-%H%M%S-%f')
    try:
        out.mkdir(parents=True,exist_ok=False)
        (out/'evidence').mkdir()
    except OSError:
        print('输出目录已存在或不可写，请选择新目录',file=sys.stderr)
        return 2
    manifest=dict(schema_version='1.0',specification_version='1.1',script_version='1.1',
                  started_at=now(),timezone='Asia/Shanghai',python=platform.python_version(),
                  config={k:getattr(a,k) for k in DEFAULTS},
                  arguments={k:str(v) if isinstance(v,Path) else v for k,v in vars(a).items()},
                  script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                  input_sha256=hashlib.sha256(('x '*a.input_units).encode()).hexdigest())
    dump(out/'manifest.json',redact(manifest,key))
    dump(out/'request_template.json',body(a,0,'x '*a.input_units))
    rows,observed,preflight=asyncio.run(run(a,key,out))
    summary=summarize(rows,a,observed)
    if not rows:
        summary.update(measurement_completed=False,stopped_reason='探测失败或缺少所需usage，未发起正式负载',
                       preflight_http=preflight['http_status'],preflight_error=preflight.get('error_response',preflight['error_type']))
    dump(out/'summary.json',summary)
    report(out,rows,summary,a)
    manifest.update(ended_at=now(),actual_requests=sum(r['started_s'] is not None for r in rows)+1)
    dump(out/'manifest.json',redact(manifest,key))
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    print(f'结果目录：{out.resolve()}')
    return 0 if summary['measurement_completed'] and summary['success_rate_passed'] else 1


if __name__=='__main__': sys.exit(main())
